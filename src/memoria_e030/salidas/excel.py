"""Compilador a Excel con fórmulas vivas, para entidades que exigen el .xlsx.

Decisiones (ver README):
- Fórmulas con nombres definidos (=Z*U*C_R*S*P), legibles por cualquier revisor.
- Sin LAMBDA: muchas entidades usan Excel 2016/2019, que no la soportan.
- Excel reserva los nombres "C" y "R" (notación F1C1), así que esos dos se
  llaman C_amp y R_coef.
- Las tablas de la norma viven dentro del libro: las búsquedas también son
  fórmulas, no números pegados.

Las fórmulas de este archivo son una segunda implementación del cálculo. Lo que
las mantiene honestas es `verificacion.py`, que recalcula el libro generado y
lo compara contra el motor en Python.
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ..calculo import Resultado
from ..entrada import Entrada
from ..norma import e030
from ..trazabilidad import Trazabilidad

GENERADOR = "memoria-e030"

AZUL = PatternFill("solid", fgColor="1F3A5F")
GRIS = PatternFill("solid", fgColor="EEF1F5")
AMARILLO = PatternFill("solid", fgColor="FFF4C2")
BLANCO = Font(color="FFFFFF", bold=True)
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14, color="1F3A5F")
FINO = Side(style="thin", color="C9D1DC")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
DECIMALES = "0.000"


@dataclass
class MapaCeldas:
    """Dónde quedó cada resultado, para que la verificación lo encuentre."""
    memoria: dict[str, str]          # atributo de Resultado -> "Hoja!F5"
    pisos: dict[str, list[str]]      # atributo de FilaPiso -> celdas por nivel


def _nombrar(wb, nombre, hoja, ref):
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"{hoja}!{ref}")


def _encabezado(ws, fila, textos):
    for col, texto in enumerate(textos, start=1):
        c = ws.cell(fila, col, texto)
        c.fill, c.font, c.border = AZUL, BLANCO, BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _anchos(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _portada(wb, entrada: Entrada, traza: Trazabilidad):
    ws = wb.active
    ws.title = "Portada"
    p = entrada.proyecto
    ws["A1"] = "MEMORIA DE CÁLCULO — ANÁLISIS SÍSMICO ESTÁTICO"
    ws["A1"].font = TITULO
    ws["A2"] = f"Norma Técnica {e030.NORMA} Diseño Sismorresistente"
    filas = [
        ("Proyecto", p.nombre),
        ("Ubicación", p.ubicacion),
        ("Entidad", p.entidad),
        ("Contratista", p.contratista),
        ("Fecha", p.fecha),
    ]
    filas += [(e.rol, f"{e.nombre} — CIP {e.cip}") for e in p.especialistas]
    filas += [
        ("", ""),
        ("Trazabilidad", ""),
        ("Versión del código", traza.version),
        ("Estado", traza.estado),
        ("Generado", traza.fecha),
        ("Generador", GENERADOR),
        ("", ""),
        ("Cómo leer este libro", "Datos: celdas amarillas editables. Memoria y Distribucion: fórmulas vivas."),
    ]
    for i, (k, v) in enumerate(filas, start=4):
        ws.cell(i, 1, k).font = NEGRITA
        ws.cell(i, 2, v)
    _anchos(ws, [26, 90])


def _tablas(wb):
    ws = wb.create_sheet("Tablas_E030")
    ws["A1"] = f"Tablas de la norma {e030.NORMA} (transcripción: normas/E030-2018/)"
    ws["A1"].font = TITULO

    def lista(fila, col, titulo, cabeceras, datos, nombres):
        ws.cell(fila, col, titulo).font = NEGRITA
        for j, c in enumerate(cabeceras):
            x = ws.cell(fila + 1, col + j, c)
            x.fill, x.font, x.border = AZUL, BLANCO, BORDE
        for i, fila_datos in enumerate(datos, start=fila + 2):
            for j, v in enumerate(fila_datos):
                ws.cell(i, col + j, v).border = BORDE
        ini, fin = fila + 2, fila + 1 + len(datos)
        for j, nombre in enumerate(nombres):
            if nombre:
                letra = get_column_letter(col + j)
                _nombrar(wb, nombre, ws.title, f"${letra}${ini}:${letra}${fin}")
        return ini, fin

    lista(3, 1, "Tabla N° 1 — Factor de zona (Art. 10)", ["Zona", "Z"],
          list(e030.FACTOR_ZONA.items()), ["Tabla1_Zona", "Tabla1_Z"])

    suelos = list(e030.PERIODOS_SUELO)
    ini, fin = lista(3, 4, "Tabla N° 3 — Factor de suelo S (Art. 13.1)", ["Zona", *suelos],
                     [[z, *(e030.FACTOR_SUELO[z][s] for s in suelos)] for z in e030.FACTOR_SUELO],
                     ["Tabla3_Zona"])
    _nombrar(wb, "Tabla3_S", ws.title, f"$E${ini}:$H${fin}")
    _nombrar(wb, "Tabla3_Suelo", ws.title, "$E$4:$H$4")

    lista(10, 1, "Tabla N° 4 — Periodos (Art. 13.2)", ["Suelo", "T_P (s)", "T_L (s)"],
          [[s, *e030.PERIODOS_SUELO[s]] for s in suelos], ["Tabla4_Suelo", "Tabla4_TP", "Tabla4_TL"])

    lista(10, 5, "Tabla N° 5 — Factor de uso (Art. 15)", ["Categoría", "U"],
          list(e030.FACTOR_USO.items()), ["Tabla5_Categoria", "Tabla5_U"])

    lista(18, 1, "Tabla N° 7 — Coeficiente básico de reducción (Art. 18)", ["Sistema", "R_0"],
          list(e030.COEF_R0.items()), ["Tabla7_Sistema", "Tabla7_R0"])

    lista(18, 5, "Fracción de CV en el peso sísmico (Art. 26)", ["Categoría", "f"],
          list(e030.FRACCION_CV.items()), ["TablaCV_Cat", "TablaCV_f"])

    ws["E24"], ws["F24"] = "Azotea", e030.FRACCION_CV_AZOTEA
    ws["E25"], ws["F25"] = "C/R mínimo (Art. 28.2.1)", e030.C_R_MINIMO
    ws["E24"].font = ws["E25"].font = NEGRITA
    _nombrar(wb, "f_azotea", ws.title, "$F$24")
    _nombrar(wb, "C_R_min", ws.title, "$F$25")
    _anchos(ws, [36, 10, 10, 10, 26, 10, 10, 10])


FILA_PISOS = 13  # primera fila de pisos en la hoja Datos


def _datos(wb, entrada: Entrada):
    ws = wb.create_sheet("Datos")
    ws["A1"] = "DATOS DE ENTRADA (celdas amarillas)"
    ws["A1"].font = TITULO
    q = entrada.parametros
    parametros = [
        ("Zona sísmica (Art. 10)", "Zona", q.zona),
        ("Perfil de suelo (Art. 12)", "Suelo", q.suelo),
        ("Categoría de la edificación (Art. 15)", "Categoria", q.categoria),
        ("Sistema estructural (Art. 18)", "Sistema", q.sistema),
        ("Irregularidad en altura I_a (Art. 20)", "I_a", q.I_a),
        ("Irregularidad en planta I_p (Art. 20)", "I_p", q.I_p),
        ("Coeficiente C_T (Art. 28.4.1)", "C_T", q.C_T),
    ]
    for i, (texto, nombre, valor) in enumerate(parametros, start=3):
        ws.cell(i, 1, texto)
        c = ws.cell(i, 2, valor)
        c.fill, c.border = AMARILLO, BORDE
        _nombrar(wb, nombre, ws.title, f"$B${i}")

    _encabezado(ws, FILA_PISOS - 1, ["Nivel", "h entrepiso (m)", "CM (tonf)", "CV (tonf)", "Azotea (SI/NO)"])
    for i, p in enumerate(entrada.pisos, start=FILA_PISOS):
        for col, v in enumerate([p.nivel, p.h, p.CM, p.CV, "SI" if p.azotea else "NO"], start=1):
            c = ws.cell(i, col, v)
            c.border = BORDE
            if col > 1:
                c.fill = AMARILLO
    _anchos(ws, [40, 18, 14, 14, 16])


# (atributo de Resultado, nombre Excel, descripción, referencia, expresión, fórmula, unidad)
PASOS = [
    ("Z", "Z", "Factor de zona", "Art. 10, Tabla N° 1", "Tabla N° 1 según zona",
     "=INDEX(Tabla1_Z,MATCH(Zona,Tabla1_Zona,0))", "—"),
    ("S", "S", "Factor de suelo", "Art. 13.1, Tabla N° 3", "Tabla N° 3 según zona y suelo",
     "=INDEX(Tabla3_S,MATCH(Zona,Tabla3_Zona,0),MATCH(Suelo,Tabla3_Suelo,0))", "—"),
    ("T_P", "T_P", "Periodo que define la plataforma", "Art. 13.2, Tabla N° 4", "Tabla N° 4 según suelo",
     "=INDEX(Tabla4_TP,MATCH(Suelo,Tabla4_Suelo,0))", "s"),
    ("T_L", "T_L", "Periodo de inicio de desplazamiento constante", "Art. 13.2, Tabla N° 4", "Tabla N° 4 según suelo",
     "=INDEX(Tabla4_TL,MATCH(Suelo,Tabla4_Suelo,0))", "s"),
    ("U", "U", "Factor de uso", "Art. 15, Tabla N° 5", "Tabla N° 5 según categoría",
     "=INDEX(Tabla5_U,MATCH(Categoria,Tabla5_Categoria,0))", "—"),
    ("R_0", "R_0", "Coeficiente básico de reducción", "Art. 18, Tabla N° 7", "Tabla N° 7 según sistema",
     "=INDEX(Tabla7_R0,MATCH(Sistema,Tabla7_Sistema,0))", "—"),
    ("R", "R_coef", "Coeficiente de reducción", "Art. 22", "R = R_0 · I_a · I_p",
     "=R_0*I_a*I_p", "—"),
    ("h_n", "h_n", "Altura total de la edificación", "Art. 28.4.1", "h_n = Σ h entrepiso",
     f"=SUM(Datos!$B${FILA_PISOS}:$B${{fin}})", "m"),
    ("T", "T", "Periodo fundamental de vibración", "Art. 28.4.1", "T = h_n / C_T",
     "=h_n/C_T", "s"),
    ("C", "C_amp", "Factor de amplificación sísmica", "Art. 14",
     "T<T_P: 2.5 | T<T_L: 2.5·T_P/T | si no: 2.5·T_P·T_L/T²",
     "=IF(T<T_P,2.5,IF(T<T_L,2.5*T_P/T,2.5*T_P*T_L/T^2))", "—"),
    ("C_R_calculado", "C_R_calc", "Cociente C/R calculado", "Art. 28.2.1", "C / R",
     "=C_amp/R_coef", "—"),
    ("C_R", "C_R", "Cociente C/R de diseño", "Art. 28.2.1", "máx(C/R ; 0.11)",
     "=MAX(C_R_calc,C_R_min)", "—"),
    ("P", "P", "Peso sísmico total", "Art. 26", "P = Σ P_i (hoja Distribucion)",
     "=SUM(Distribucion!$F$4:$F${fin_d})", "tonf"),
    ("V", "V", "Fuerza cortante en la base", "Art. 28.2.1", "V = Z · U · (C/R) · S · P",
     "=Z*U*C_R*S*P", "tonf"),
    ("k", "k", "Exponente de distribución en altura", "Art. 28.3.2",
     "T≤0.5: 1 | si no: mín(0.75+0.5·T ; 2)",
     "=IF(T<=0.5,1,MIN(0.75+0.5*T,2))", "—"),
]


def _memoria(wb, n_pisos) -> dict[str, str]:
    ws = wb.create_sheet("Memoria")
    ws["A1"] = "MEMORIA — ANÁLISIS ESTÁTICO (E.030 Art. 28)"
    ws["A1"].font = TITULO
    _encabezado(ws, 3, ["N°", "Descripción", "Referencia", "Símbolo", "Expresión", "Valor", "Unidad"])
    fin = FILA_PISOS + n_pisos - 1
    fin_d = 3 + n_pisos
    mapa = {}
    for i, (attr, nombre, desc, ref, expr, formula, unidad) in enumerate(PASOS, start=4):
        formula = formula.replace("{fin}", str(fin)).replace("{fin_d}", str(fin_d))
        for col, v in enumerate([i - 3, desc, ref, nombre, expr, formula, unidad], start=1):
            c = ws.cell(i, col, v)
            c.border = BORDE
            if i % 2:
                c.fill = GRIS
        ws.cell(i, 6).number_format = DECIMALES
        ws.cell(i, 6).font = NEGRITA
        _nombrar(wb, nombre, ws.title, f"$F${i}")
        mapa[attr] = f"Memoria!F{i}"
    ws.freeze_panes = "A4"
    _anchos(ws, [5, 44, 22, 11, 52, 12, 8])
    return mapa


def _distribucion(wb, n_pisos) -> dict[str, list[str]]:
    ws = wb.create_sheet("Distribucion")
    ws["A1"] = "DISTRIBUCIÓN DE FUERZAS EN ALTURA (E.030 Art. 28.3)"
    ws["A1"].font = TITULO
    _encabezado(ws, 3, ["Nivel", "h_i (m)", "CM (tonf)", "CV (tonf)", "f", "P_i (tonf)",
                        "P_i·h_i^k", "α_i", "F_i (tonf)", "V_i (tonf)"])
    ini, fin = 4, 3 + n_pisos
    mapa = {"h_i": [], "f": [], "P_i": [], "alfa": [], "F_i": [], "V_i": []}
    for j in range(n_pisos):
        r, d = ini + j, FILA_PISOS + j
        formulas = [
            f"=Datos!A{d}",
            f"=SUM(Datos!$B${FILA_PISOS}:B{d})",
            f"=Datos!C{d}",
            f"=Datos!D{d}",
            f'=IF(Datos!E{d}="SI",f_azotea,INDEX(TablaCV_f,MATCH(LEFT(Categoria,1),TablaCV_Cat,0)))',
            f"=C{r}+E{r}*D{r}",
            f"=F{r}*B{r}^k",
            f"=G{r}/SUM($G${ini}:$G${fin})",
            f"=H{r}*V",
            f"=SUM(I{r}:$I${fin})",
        ]
        for col, f in enumerate(formulas, start=1):
            c = ws.cell(r, col, f)
            c.border = BORDE
            if col > 1:
                c.number_format = DECIMALES
        for attr, col in [("h_i", "B"), ("f", "E"), ("P_i", "F"), ("alfa", "H"), ("F_i", "I"), ("V_i", "J")]:
            mapa[attr].append(f"Distribucion!{col}{r}")
    t = fin + 1
    ws.cell(t, 1, "Σ").font = NEGRITA
    for col in "FHI":
        c = ws[f"{col}{t}"]
        c.value = f"=SUM({col}{ini}:{col}{fin})"
        c.font, c.number_format = NEGRITA, DECIMALES
    ws.freeze_panes = "A4"
    _anchos(ws, [8, 10, 11, 11, 7, 12, 13, 9, 12, 12])
    return mapa


def compilar(entrada: Entrada, resultado: Resultado, traza: Trazabilidad, destino: Path) -> MapaCeldas:
    n = len(entrada.pisos)
    wb = Workbook()
    _portada(wb, entrada, traza)
    _datos(wb, entrada)
    memoria = _memoria(wb, n)
    pisos = _distribucion(wb, n)
    _tablas(wb)
    # Metadatos del archivo: nunca el usuario del sistema operativo
    wb.properties.creator = GENERADOR
    wb.properties.lastModifiedBy = GENERADOR
    wb.properties.title = f"Memoria E.030 — {entrada.proyecto.nombre}"
    wb.active = wb.sheetnames.index("Memoria")
    wb.save(destino)
    return MapaCeldas(memoria, pisos)
